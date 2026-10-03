import csv
import json
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from test_export_ledger_csv import HEADERS, ROOT, complete_row, page


SCRIPT = ROOT / "scripts" / "export_ledger.py"


class ExportLedgerXlsxTests(unittest.TestCase):
    def run_export(self, directory, rows, *arguments):
        source = directory / "input.json"
        source.write_text(json.dumps(page(rows), ensure_ascii=False), encoding="utf-8")
        return subprocess.run(
            [sys.executable, str(SCRIPT), str(source), *map(str, arguments)],
            capture_output=True, text=True,
        )

    def test_xlsx_keeps_phones_identifiers_and_formula_like_values_as_literal_text(self):
        with tempfile.TemporaryDirectory() as tmp:
            directory = Path(tmp)
            target = directory / "ledger.xlsx"
            result = self.run_export(directory, [
                complete_row(supplier_contact_phone="01012345678", receipt_id="0000123", item_name="=1+1"),
                complete_row(supplier_contact_phone="+8613800138000", specification="@规格"),
                complete_row(supplier_contact_phone=None),
            ], target, "--delete-input")
            self.assertEqual(result.returncode, 0, result.stderr)
            self.assertFalse((directory / "input.json").exists())
            book = load_workbook(target)
            sheet = book.active
            self.assertEqual([cell.value for cell in sheet[1]], HEADERS)
            self.assertEqual(sheet.max_column, 11)
            for address, value in {"A2": "0000123", "B2": "=1+1", "J2": "01012345678", "J3": "+8613800138000", "C3": "@规格"}.items():
                self.assertEqual(sheet[address].value, value)
                self.assertEqual(sheet[address].data_type, "s")
                self.assertEqual(sheet[address].number_format, "@")
            self.assertIsNone(sheet["J4"].value)
            self.assertEqual(sheet["D2"].value, 12.5)
            self.assertEqual(sheet["E2"].value, 99.9)
            self.assertEqual(sheet["D2"].data_type, "n")
            book.close()

    def test_default_monthly_xlsx_skips_existing_receipts_and_preserves_text_on_append(self):
        with tempfile.TemporaryDirectory() as tmp:
            directory = Path(tmp)
            args = ("--store-month-dir", directory, "--month", "2026-09")
            old = complete_row(store_name="通州店", receipt_id="OLD", supplier_contact_phone="01012345678")
            result = self.run_export(directory, [old], *args)
            self.assertEqual(result.returncode, 0, result.stderr)
            target = directory / "食品经营单位进货台帐_通州店_2026-09.xlsx"
            self.assertTrue(target.exists())
            existing_bytes = target.read_bytes()
            result = self.run_export(directory, [old], *args)
            self.assertEqual(result.returncode, 0, result.stderr)
            self.assertEqual(target.read_bytes(), existing_bytes)
            self.assertEqual(json.loads(result.stdout)["receipts_skipped"], 1)
            new = complete_row(store_name="通州店", receipt_id="NEW", supplier_contact_phone="15012345678")
            result = self.run_export(directory, [old, new, new], *args)
            self.assertEqual(result.returncode, 0, result.stderr)
            book = load_workbook(target)
            self.assertEqual([row[0] for row in list(book.active.values)[1:]], ["OLD", "NEW", "NEW"])
            self.assertEqual(book.active["J2"].value, "01012345678")
            self.assertEqual(book.active["J3"].value, "15012345678")
            self.assertEqual(book.active["J3"].data_type, "s")
            book.close()

    def test_explicit_csv_remains_available_without_creating_xlsx(self):
        with tempfile.TemporaryDirectory() as tmp:
            directory = Path(tmp)
            target = directory / "ledger.csv"
            result = self.run_export(directory, [complete_row(supplier_contact_phone="01012345678")], target)
            self.assertEqual(result.returncode, 0, result.stderr)
            with target.open(encoding="utf-8-sig", newline="") as handle:
                self.assertEqual(list(csv.reader(handle))[1][9], "01012345678")
            result = self.run_export(directory, [complete_row(store_name="通州店")], "--store-month-dir", directory, "--month", "2026-09", "--format", "csv")
            self.assertEqual(result.returncode, 0, result.stderr)
            self.assertTrue((directory / "食品经营单位进货台帐_通州店_2026-09.csv").exists())
            self.assertEqual(list(directory.glob("*.xlsx")), [])

    def test_rejects_numeric_and_scientific_notation_phones_before_overwrite_in_both_formats(self):
        for suffix in ("xlsx", "csv"):
            for phone in (15012345678, "1.501E+10"):
                with self.subTest(suffix=suffix, phone=phone), tempfile.TemporaryDirectory() as tmp:
                    directory = Path(tmp)
                    target = directory / f"ledger.{suffix}"
                    target.write_bytes(b"existing file")
                    result = self.run_export(directory, [complete_row(supplier_contact_phone=phone)], target, "--overwrite", "--delete-input")
                    self.assertNotEqual(result.returncode, 0)
                    self.assertIn("supplier_contact_phone", result.stderr)
                    self.assertEqual(target.read_bytes(), b"existing file")
                    self.assertFalse((directory / "input.json").exists())

    def test_monthly_validation_rejects_existing_numeric_phone_before_writing_any_store(self):
        with tempfile.TemporaryDirectory() as tmp:
            directory = Path(tmp)
            args = ("--store-month-dir", directory, "--month", "2026-09")
            result = self.run_export(directory, [complete_row(store_name="通州店", receipt_id="OLD")], *args)
            self.assertEqual(result.returncode, 0, result.stderr)
            target = directory / "食品经营单位进货台帐_通州店_2026-09.xlsx"
            book = load_workbook(target)
            book.active["J2"] = 15012345678
            book.save(target)
            book.close()
            original = target.read_bytes()
            result = self.run_export(directory, [complete_row(store_name="海淀店"), complete_row(store_name="通州店", receipt_id="NEW")], *args)
            self.assertNotEqual(result.returncode, 0)
            self.assertIn("text", result.stderr)
            self.assertEqual(target.read_bytes(), original)
            self.assertFalse((directory / "食品经营单位进货台帐_海淀店_2026-09.xlsx").exists())

    def test_empty_xlsx_result_leaves_existing_output_byte_identical(self):
        with tempfile.TemporaryDirectory() as tmp:
            directory = Path(tmp)
            target = directory / "ledger.xlsx"
            target.write_bytes(b"existing file")
            result = self.run_export(directory, [], target, "--overwrite")
            self.assertEqual(result.returncode, 0, result.stderr)
            self.assertEqual(target.read_bytes(), b"existing file")
            self.assertEqual(json.loads(result.stdout)["files_written"], 0)

    def test_xlsx_overwrite_is_explicit_and_format_mismatch_fails_before_write(self):
        with tempfile.TemporaryDirectory() as tmp:
            directory = Path(tmp)
            target = directory / "ledger.xlsx"
            target.write_bytes(b"existing file")
            result = self.run_export(directory, [complete_row()], target)
            self.assertNotEqual(result.returncode, 0)
            self.assertEqual(target.read_bytes(), b"existing file")
            result = self.run_export(directory, [complete_row()], target, "--format", "csv", "--overwrite")
            self.assertNotEqual(result.returncode, 0)
            self.assertEqual(target.read_bytes(), b"existing file")
            result = self.run_export(directory, [complete_row()], target, "--overwrite")
            self.assertEqual(result.returncode, 0, result.stderr)
            book = load_workbook(target)
            self.assertEqual(book.active["J2"].value, "+8613800138000")
            book.close()
            self.assertEqual(list(directory.glob("*.tmp")), [])


if __name__ == "__main__":
    unittest.main()
