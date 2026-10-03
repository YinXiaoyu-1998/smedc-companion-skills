#!/usr/bin/env python3
"""Shared ledger exporter; this legacy entry point still defaults to CSV."""

from __future__ import annotations

import argparse
import csv
import datetime as dt
import json
import os
import re
import sys
import tempfile
import unicodedata
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any


ROOT = Path(__file__).resolve().parents[1]
PROFILE_PATH = ROOT / "config" / "food-purchase-ledger-cn-v2.json"
DATASET = "delivery_ledger"
MODE = "detail"
NULLABLE_FIELDS = {
    "production_date_or_batch",
    "shelf_life",
    "supplier_unit_address",
    "supplier_contact_phone",
}
FORMULA_PREFIXES = ("=", "+", "-", "@")
DATE_RE = re.compile(r"^\d{4}-\d{2}-\d{2}$")
MONTH_RE = re.compile(r"^\d{4}-\d{2}$")
INVALID_STORE_NAME_RE = re.compile(r'[\x00-\x1f<>:"/\\|?*]')
STORE_NAME_FIELD = "store_name"
NUMERIC_FIELDS = {"purchase_quantity", "purchase_amount"}
SCIENTIFIC_RE = re.compile(r"^[+-]?(?:\d+(?:\.\d*)?|\.\d+)[eE][+-]?\d+$")


class ExportError(ValueError):
    pass


def load_expected_profile() -> dict[str, Any]:
    with PROFILE_PATH.open("r", encoding="utf-8") as handle:
        profile = json.load(handle)
    if not isinstance(profile, dict):
        raise ExportError("local profile config is malformed")
    return profile


EXPECTED_PROFILE = load_expected_profile()
CANONICAL_FIELDS = [column["canonicalName"] for column in EXPECTED_PROFILE["columns"]]
HEADERS = [column["displayName"] for column in EXPECTED_PROFILE["columns"]]
RECEIPT_INDEX = CANONICAL_FIELDS.index("receipt_id")
PURCHASE_DATE_INDEX = CANONICAL_FIELDS.index("purchase_date")


class NormalizedRow:
    def __init__(self, cells: list[str], store_name: str | None, page_index: int, row_index: int) -> None:
        self.cells = cells
        self.store_name = store_name
        self.page_index = page_index
        self.row_index = row_index

    @property
    def receipt_id(self) -> str:
        return self.cells[RECEIPT_INDEX]

    @property
    def purchase_date(self) -> str:
        return self.cells[PURCHASE_DATE_INDEX]


def parse_args(argv: list[str], default_format: str) -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Export SMEDC delivery ledger detail results to XLSX or CSV.")
    parser.add_argument("input_json", type=Path)
    parser.add_argument("output_csv", type=Path, nargs="?", metavar="OUTPUT_FILE")
    parser.add_argument("--overwrite", action="store_true", help="replace an existing output file")
    parser.add_argument("--store-month-dir", type=Path, help="create or maintain per-store monthly ledger files in this directory")
    parser.add_argument("--month", help="requested month for --store-month-dir, formatted YYYY-MM")
    parser.add_argument("--format", choices=("xlsx", "csv"), help="output format; inferred from the output suffix, otherwise the entry point default")
    parser.add_argument(
        "--delete-input",
        action="store_true",
        help="delete the agent-created input JSON after this run, including validation failures",
    )
    args = parser.parse_args(argv)
    suffix = args.output_csv.suffix.lower().lstrip(".") if args.output_csv else None
    args.format = args.format or (suffix if suffix in ("xlsx", "csv") else default_format)
    return args


def load_payload(path: Path) -> Any:
    try:
        with path.open("r", encoding="utf-8") as handle:
            return json.load(handle)
    except OSError as exc:
        raise ExportError(f"unable to read input JSON: {exc}") from exc
    except json.JSONDecodeError as exc:
        raise ExportError(f"input JSON is malformed: {exc}") from exc


def as_pages(payload: Any) -> list[dict[str, Any]]:
    if isinstance(payload, dict):
        return [payload]
    if isinstance(payload, list) and all(isinstance(page, dict) for page in payload):
        return payload
    raise ExportError("input must be one page object or an array of page objects")


def page_fingerprint(page: dict[str, Any]) -> str:
    return json.dumps(page, ensure_ascii=False, sort_keys=True, separators=(",", ":"))


def page_profile(page: dict[str, Any]) -> Any:
    if "presentation" in page:
        return page["presentation"]
    raise ExportError("page is missing presentation")


def page_rows(page: dict[str, Any]) -> Any:
    for key in ("rows", "records", "data"):
        if key in page:
            return page[key]
    raise ExportError("page is missing rows")


def validate_date(value: str, field: str) -> None:
    if not DATE_RE.fullmatch(value):
        raise ExportError(f"{field} must be YYYY-MM-DD")
    try:
        dt.date.fromisoformat(value)
    except ValueError as exc:
        raise ExportError(f"{field} must be a valid YYYY-MM-DD date") from exc


def validate_month(value: str) -> None:
    if not MONTH_RE.fullmatch(value):
        raise ExportError("month must be YYYY-MM")
    try:
        dt.date.fromisoformat(f"{value}-01")
    except ValueError as exc:
        raise ExportError("month must be a valid YYYY-MM month") from exc


def scalar_to_string(value: Any, field: str, page_index: int, row_index: int) -> str:
    if value is None:
        raise ExportError(f"row {page_index}.{row_index} field {field} cannot be null")
    if isinstance(value, (list, dict)):
        raise ExportError(f"row {page_index}.{row_index} field {field} must be scalar")
    if isinstance(value, bool):
        return "true" if value else "false"
    return str(value)


def normalize_cell(row: dict[str, Any], field: str, page_index: int, row_index: int, *, csv_safe: bool = True) -> str:
    if field not in row:
        if field in NULLABLE_FIELDS:
            return ""
        raise ExportError(f"row {page_index}.{row_index} is missing required field {field}")

    value = row[field]
    if value is None:
        if field in NULLABLE_FIELDS:
            return ""
        raise ExportError(f"row {page_index}.{row_index} field {field} cannot be null")
    if field in NULLABLE_FIELDS and value == "null":
        return ""
    if field == "supplier_contact_phone":
        if not isinstance(value, str):
            raise ExportError(f"row {page_index}.{row_index} supplier_contact_phone must be a string or null")
        if SCIENTIFIC_RE.fullmatch(value.strip()):
            raise ExportError(f"row {page_index}.{row_index} supplier_contact_phone must not use scientific notation; obtain the original phone string")
    cell = scalar_to_string(value, field, page_index, row_index)
    if not csv_safe and field in NUMERIC_FIELDS:
        xlsx_number(cell, field)

    if field == "purchase_date":
        validate_date(cell, field)
    if field == "receipt_id" and not cell:
        raise ExportError(f"row {page_index}.{row_index} field receipt_id cannot be empty")

    if csv_safe and cell.startswith(FORMULA_PREFIXES):
        return "'" + cell
    return cell


def normalize_store_name(row: dict[str, Any], page_index: int, row_index: int) -> str:
    if STORE_NAME_FIELD not in row:
        raise ExportError(f"row {page_index}.{row_index} is missing required field {STORE_NAME_FIELD}")
    store_name = scalar_to_string(row[STORE_NAME_FIELD], STORE_NAME_FIELD, page_index, row_index)
    if not store_name or store_name != store_name.strip():
        raise ExportError(f"row {page_index}.{row_index} field {STORE_NAME_FIELD} must be non-empty without surrounding whitespace")
    if INVALID_STORE_NAME_RE.search(store_name):
        raise ExportError(f"row {page_index}.{row_index} field {STORE_NAME_FIELD} contains characters invalid for filenames")
    return store_name


def validate_pages(pages: list[dict[str, Any]], *, require_store_name: bool = False, csv_safe: bool = True) -> list[NormalizedRow]:
    rows_out: list[NormalizedRow] = []
    seen_pages: set[str] = set()

    for page_index, page in enumerate(pages, start=1):
        fingerprint = page_fingerprint(page)
        if fingerprint in seen_pages:
            raise ExportError(f"duplicate page payload detected at page {page_index}")
        seen_pages.add(fingerprint)

        if page.get("dataset") != DATASET:
            raise ExportError(f"page {page_index} dataset must be {DATASET}")
        if page.get("mode") != MODE:
            raise ExportError(f"page {page_index} mode must be {MODE}")
        if page_profile(page) != EXPECTED_PROFILE:
            raise ExportError(f"page {page_index} presentation does not match {EXPECTED_PROFILE['id']}")

        rows = page_rows(page)
        if not isinstance(rows, list):
            raise ExportError(f"page {page_index} rows must be an array")
        for row_index, row in enumerate(rows, start=1):
            if not isinstance(row, dict):
                raise ExportError(f"row {page_index}.{row_index} must be an object")
            cells = [
                normalize_cell(row, field, page_index, row_index, csv_safe=csv_safe)
                for field in CANONICAL_FIELDS
            ]
            store_name = normalize_store_name(row, page_index, row_index) if require_store_name else None
            rows_out.append(NormalizedRow(cells, store_name, page_index, row_index))

    return rows_out


def validate_existing_rows(path: Path, month: str) -> tuple[list[list[str]], set[str]]:
    try:
        with path.open("r", encoding="utf-8-sig", newline="") as handle:
            rows = list(csv.reader(handle))
    except OSError as exc:
        raise ExportError(f"unable to read existing CSV {path}: {exc}") from exc

    if not rows or rows[0] != HEADERS:
        raise ExportError(f"existing CSV header does not match {EXPECTED_PROFILE['id']}: {path}")

    existing_receipts: set[str] = set()
    body = rows[1:]
    for row_number, row in enumerate(body, start=2):
        if len(row) != len(HEADERS):
            raise ExportError(f"existing CSV row {row_number} has {len(row)} columns, expected {len(HEADERS)}: {path}")
        receipt_id = row[RECEIPT_INDEX]
        if not receipt_id:
            raise ExportError(f"existing CSV row {row_number} receipt_id cannot be empty: {path}")
        purchase_date = row[PURCHASE_DATE_INDEX]
        validate_date(purchase_date, f"existing CSV row {row_number} purchase_date")
        if not purchase_date.startswith(f"{month}-"):
            raise ExportError(f"existing CSV row {row_number} is outside requested month {month}: {path}")
        existing_receipts.add(receipt_id)

    return body, existing_receipts


def monthly_filename(store_name: str, month: str, file_format: str = "csv") -> str:
    return f"{EXPECTED_PROFILE['title']}_{store_name}_{month}.{file_format}"


def normalized_target_key(path: Path) -> str:
    return unicodedata.normalize("NFC", path.name).casefold()


def build_store_month_plans(rows: list[NormalizedRow], output_dir: Path, month: str, file_format: str = "csv") -> tuple[dict[Path, list[list[str]]], dict[str, Any]]:
    validate_month(month)
    if not output_dir.exists() or not output_dir.is_dir():
        raise ExportError(f"store-month output directory does not exist: {output_dir}")
    if not rows:
        return {}, {
            "mode": "store-month",
            "month": month,
            "files_written": 0,
            "receipts_appended": 0,
            "receipts_skipped": 0,
            "rows_appended": 0,
        }

    receipt_store: dict[str, str] = {}
    grouped: dict[Path, dict[str, list[NormalizedRow]]] = {}
    normalized_targets: dict[str, tuple[str, Path]] = {}
    for row in rows:
        if row.store_name is None:
            raise ExportError(f"row {row.page_index}.{row.row_index} is missing store_name")
        if not row.purchase_date.startswith(f"{month}-"):
            raise ExportError(f"row {row.page_index}.{row.row_index} purchase_date is outside requested month {month}")
        prior_store = receipt_store.setdefault(row.receipt_id, row.store_name)
        if prior_store != row.store_name:
            raise ExportError(f"receipt_id {row.receipt_id} appears under multiple stores")
        target = output_dir / monthly_filename(row.store_name, month, file_format)
        if target.name != monthly_filename(row.store_name, month, file_format):
            raise ExportError(f"invalid filename for store {row.store_name}")
        normalized_key = normalized_target_key(target)
        existing_target = normalized_targets.setdefault(normalized_key, (row.store_name, target))
        if existing_target[0] != row.store_name:
            raise ExportError(
                "filename collision after normalization: "
                f"{existing_target[0]} -> {existing_target[1].name}, {row.store_name} -> {target.name}"
            )
        grouped.setdefault(target, {}).setdefault(row.receipt_id, []).append(row)

    plans: dict[Path, list[list[str]]] = {}
    receipts_appended = 0
    receipts_skipped = 0
    rows_appended = 0

    for target, receipt_groups in grouped.items():
        if target.exists():
            existing_rows, existing_receipts = (
                validate_existing_xlsx_rows(target, month)
                if file_format == "xlsx" else validate_existing_rows(target, month)
            )
        else:
            existing_rows = []
            existing_receipts = set()

        output_rows = list(existing_rows)
        file_rows_appended = 0
        for receipt_id, receipt_rows in receipt_groups.items():
            if receipt_id in existing_receipts:
                receipts_skipped += 1
                continue
            rows_to_append = [row.cells for row in receipt_rows]
            output_rows.extend(rows_to_append)
            receipts_appended += 1
            file_rows_appended += len(rows_to_append)
            rows_appended += len(rows_to_append)

        if file_rows_appended:
            plans[target] = output_rows

    summary = {
        "mode": "store-month",
        "month": month,
        "files_written": len(plans),
        "receipts_appended": receipts_appended,
        "receipts_skipped": receipts_skipped,
        "rows_appended": rows_appended,
    }
    return plans, summary


def write_store_month_files(plans: dict[Path, list[list[str]]], file_format: str) -> None:
    for target, rows in plans.items():
        write_ledger_atomic(target, rows, overwrite=True, file_format=file_format)


def xlsx_number(value: str, field: str) -> Decimal:
    try:
        number = Decimal(value)
    except InvalidOperation as exc:
        raise ExportError(f"{field} must be numeric for XLSX export") from exc
    if not number.is_finite():
        raise ExportError(f"{field} must be finite for XLSX export")
    return number


def xlsx_library():
    try:
        import openpyxl
    except ImportError as exc:
        raise ExportError("XLSX export requires openpyxl; install this skill's requirements.txt in your Python environment") from exc
    return openpyxl


def validate_existing_xlsx_rows(path: Path, month: str) -> tuple[list[list[str]], set[str]]:
    library = xlsx_library()
    try:
        book = library.load_workbook(path, read_only=True, data_only=False, keep_links=False)
    except Exception as exc:
        raise ExportError(f"unable to read existing XLSX: {path}") from exc
    try:
        if len(book.worksheets) != 1:
            raise ExportError(f"existing XLSX must contain one ledger worksheet: {path}")
        sheet = book.active
        if sheet.max_column != len(HEADERS) or [cell.value for cell in next(sheet.iter_rows(max_row=1))] != HEADERS:
            raise ExportError(f"existing XLSX header does not match {EXPECTED_PROFILE['id']}: {path}")
        body = []
        receipts = set()
        for row_number, cells in enumerate(sheet.iter_rows(min_row=2), start=2):
            if all(cell.value is None for cell in cells):
                continue
            for field, cell in zip(CANONICAL_FIELDS, cells):
                if cell.data_type == "f" or (field not in NUMERIC_FIELDS and cell.value is not None and (not isinstance(cell.value, str) or cell.data_type not in ("s", "inlineStr"))):
                    raise ExportError(f"existing XLSX row {row_number} {field} must be literal text: {path}")
            values = dict(zip(CANONICAL_FIELDS, (cell.value for cell in cells)))
            row = [normalize_cell(values, field, 0, row_number, csv_safe=False) for field in CANONICAL_FIELDS]
            if not row[PURCHASE_DATE_INDEX].startswith(f"{month}-"):
                raise ExportError(f"existing XLSX row {row_number} is outside requested month {month}: {path}")
            body.append(row)
            receipts.add(row[RECEIPT_INDEX])
        return body, receipts
    finally:
        book.close()


def write_ledger_atomic(output_path: Path, rows: list[list[str]], overwrite: bool, file_format: str) -> None:
    if file_format == "csv":
        write_csv_atomic(output_path, rows, overwrite)
        return
    if output_path.exists() and not overwrite:
        raise ExportError(f"output already exists: {output_path}")
    if not output_path.parent.is_dir():
        raise ExportError(f"output directory does not exist: {output_path.parent}")
    library = xlsx_library()
    book = library.Workbook()
    sheet = book.active
    sheet.title = EXPECTED_PROFILE["title"]
    sheet.freeze_panes = "A2"
    for row_number, values in enumerate([HEADERS, *rows], start=1):
        for index, value in enumerate(values, start=1):
            cell = sheet.cell(row_number, index)
            field = CANONICAL_FIELDS[index - 1]
            if row_number > 1 and field in NUMERIC_FIELDS:
                cell.value = xlsx_number(value, field)
                cell.number_format = "0.00##" if field == "purchase_amount" else "General"
            else:
                cell.value = value
                cell.data_type = "s"
                cell.number_format = "@"
            if row_number == 1:
                cell.font = library.styles.Font(bold=True)
    for index, width in enumerate((24, 24, 18, 12, 14, 24, 14, 36, 36, 24, 16), start=1):
        sheet.column_dimensions[library.utils.get_column_letter(index)].width = width
    sheet.auto_filter.ref = sheet.dimensions
    temp_name = None
    try:
        with tempfile.NamedTemporaryFile(dir=output_path.parent, prefix=f".{output_path.name}.", suffix=".tmp", delete=False) as handle:
            temp_name = handle.name
        book.save(temp_name)
        os.replace(temp_name, output_path)
        temp_name = None
    finally:
        book.close()
        if temp_name:
            Path(temp_name).unlink(missing_ok=True)


def write_csv_atomic(output_path: Path, rows: list[list[str]], overwrite: bool) -> None:
    if output_path.exists() and not overwrite:
        raise ExportError(f"output already exists: {output_path}")
    parent = output_path.parent
    if not parent.exists():
        raise ExportError(f"output directory does not exist: {parent}")

    temp_name: str | None = None
    try:
        with tempfile.NamedTemporaryFile(
            "w",
            encoding="utf-8-sig",
            newline="",
            dir=parent,
            prefix=f".{output_path.name}.",
            suffix=".tmp",
            delete=False,
        ) as handle:
            temp_name = handle.name
            writer = csv.writer(handle)
            writer.writerow(HEADERS)
            writer.writerows(rows)
        os.replace(temp_name, output_path)
        temp_name = None
    finally:
        if temp_name:
            try:
                os.unlink(temp_name)
            except FileNotFoundError:
                pass


def run_export(args: argparse.Namespace) -> None:
    if (
        args.delete_input
        and args.output_csv is not None
        and args.input_json.resolve() == args.output_csv.resolve()
    ):
        raise ExportError("input JSON and output CSV must be different files")
    if args.store_month_dir:
        if args.output_csv is not None:
            raise ExportError("output_csv is not used with --store-month-dir")
        if args.overwrite:
            raise ExportError("--overwrite is not used with --store-month-dir")
        if not args.month:
            raise ExportError("--month is required with --store-month-dir")
    elif args.output_csv is None:
        raise ExportError("output_csv is required unless --store-month-dir is used")
    elif args.month:
        raise ExportError("--month is only used with --store-month-dir")
    if args.output_csv is not None and args.output_csv.suffix.lower() != f".{args.format}":
        raise ExportError(f"output filename must end in .{args.format}")

    payload = load_payload(args.input_json)
    pages = as_pages(payload)
    if args.store_month_dir:
        rows = validate_pages(pages, require_store_name=True, csv_safe=args.format == "csv")
        plans, summary = build_store_month_plans(rows, args.store_month_dir, args.month, args.format)
        write_store_month_files(plans, args.format)
        print(json.dumps(summary, ensure_ascii=False, sort_keys=True))
    else:
        rows = validate_pages(pages, csv_safe=args.format == "csv")
        if not rows:
            print(json.dumps({
                "files_written": 0,
                "mode": "single-file",
                "rows_exported": 0,
            }, ensure_ascii=False, sort_keys=True))
            return
        write_ledger_atomic(args.output_csv, [row.cells for row in rows], args.overwrite, args.format)


def main(argv: list[str], *, default_format: str = "csv") -> int:
    args = parse_args(argv, default_format)
    exit_code = 0
    try:
        run_export(args)
    except ExportError as exc:
        print(f"error: {exc}", file=sys.stderr)
        exit_code = 1
    finally:
        if args.delete_input:
            try:
                args.input_json.unlink()
            except FileNotFoundError:
                pass
            except OSError as exc:
                print(f"error: unable to delete input JSON: {exc}", file=sys.stderr)
                exit_code = 1
    return exit_code


if __name__ == "__main__":
    raise SystemExit(main(sys.argv[1:]))
